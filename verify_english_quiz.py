import re
from pathlib import Path
from docx import Document

path = Path(r"C:\Users\1090602\Desktop\歷年檔案存放\課堂及時問答\英文基礎文法選擇題_50題.docx")
doc = Document(path)
paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

questions = []
for idx, text in enumerate(paras):
    match = re.match(r"^(\d+)\.\s", text)
    if match:
        number = int(match.group(1))
        if 1 <= number <= 50:
            questions.append((number, text, idx))

numbers = [q[0] for q in questions]
assert numbers == list(range(1, 51)), numbers
stems = [re.sub(r"^\d+\.\s*", "", q[1]) for q in questions]
assert len(set(stems)) == 50, "發現重複題幹"

for number, _, idx in questions:
    assert idx + 1 < len(paras), f"第 {number} 題缺少選項"
    choices = paras[idx + 1]
    for label in ("A.", "B.", "C.", "D."):
        assert label in choices, f"第 {number} 題缺少 {label} 選項"

assert len(doc.tables) == 2, f"預期 2 個表格，實際 {len(doc.tables)}"
answer_table = doc.tables[-1]
assert len(answer_table.rows) == 6 and len(answer_table.columns) == 10
answer_cells = [cell.text.strip() for row in answer_table.rows[1:] for cell in row.cells]
assert len(answer_cells) == 50
for number, cell_text in enumerate(answer_cells, 1):
    assert re.fullmatch(fr"{number}\. [ABCD]", cell_text), cell_text

print(f"PASS: {path.name}")
print("題目：50 題，題號 1-50，無重複題幹")
print("選項：每題均含 A、B、C、D")
print("答案：50 題，格式完整")
print(f"檔案大小：{path.stat().st_size} bytes")
