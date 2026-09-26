from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"C:\Users\1090602\Desktop\歷年檔案存放\課堂及時問答\英文基礎文法選擇題_50題.docx")


SECTIONS = [
    ("第一部分  疑問詞翻譯", "請選出正確的中文或英文意思。", [
        (1, "what 的中文意思是什麼？", ["誰", "哪裡", "什麼", "何時"], "C"),
        (2, "who 的中文意思是什麼？", ["為什麼", "誰", "如何", "哪裡"], "B"),
        (3, "where 的中文意思是什麼？", ["何時", "什麼", "誰", "哪裡"], "D"),
        (4, "when 的中文意思是什麼？", ["何時", "如何", "為什麼", "什麼"], "A"),
        (5, "why 的中文意思是什麼？", ["哪裡", "誰", "為什麼", "何時"], "C"),
        (6, "how 的中文意思是什麼？", ["什麼", "如何", "誰", "哪裡"], "B"),
        (7, "「什麼」的英文是哪一個？", ["what", "who", "when", "why"], "A"),
        (8, "「誰」的英文是哪一個？", ["where", "why", "who", "how"], "C"),
        (9, "「哪裡」的英文是哪一個？", ["when", "where", "what", "who"], "B"),
        (10, "「何時」的英文是哪一個？", ["why", "how", "where", "when"], "D"),
        (11, "「為什麼」的英文是哪一個？", ["why", "who", "what", "when"], "A"),
        (12, "「如何／怎麼樣」的英文是哪一個？", ["when", "where", "how", "why"], "C"),
    ]),
    ("第二部分  名詞複數", "請選出正確的複數名詞。", [
        (13, "book 的複數是哪一個？", ["bookes", "books", "book", "bookies"], "B"),
        (14, "pen 的複數是哪一個？", ["pens", "penes", "pen", "penses"], "A"),
        (15, "cup 的複數是哪一個？", ["cupps", "cupes", "cups", "cup"], "C"),
        (16, "pencil 的複數是哪一個？", ["penciles", "pencil", "pencilses", "pencils"], "D"),
        (17, "watch 的複數是哪一個？", ["watchs", "watches", "watch", "watchies"], "B"),
        (18, "class 的複數是哪一個？", ["classes", "classs", "class", "classies"], "A"),
        (19, "box 的複數是哪一個？", ["boxs", "box", "boxes", "boxies"], "C"),
        (20, "brush 的複數是哪一個？", ["brushs", "brush", "brushies", "brushes"], "D"),
    ]),
    ("第三部分  限定詞與單複數", "請選出最適合填入空格的答案。", [
        (21, "This is a ______.", ["books", "bookes", "book", "booking"], "C"),
        (22, "These are my ______.", ["pencils", "pencil", "penciles", "a pencil"], "A"),
        (23, "That is ______ eraser.", ["a", "two", "these", "an"], "D"),
        (24, "Those are three ______.", ["box", "boxes", "boxs", "a box"], "B"),
        (25, "I have two ______.", ["watches", "watch", "watchs", "a watch"], "A"),
        (26, "______ cup is on the desk.", ["These", "Those", "This", "They"], "C"),
        (27, "______ books are interesting.", ["That", "Those", "This", "It"], "B"),
        (28, "She has ______ orange.", ["a", "two", "this", "an"], "D"),
    ]),
    ("第四部分  以 What 詢問物品", "請選出最適合的答案。", [
        (29, "A: What is this?\nB: ______ a pencil.", ["They’re", "It’s", "We’re", "You’re"], "B"),
        (30, "A: What is that?\nB: It’s ______ eraser.", ["an", "a", "two", "these"], "A"),
        (31, "A: What are these?\nB: ______ my books.", ["It’s", "He’s", "They’re", "She’s"], "C"),
        (32, "A: What are those?\nB: They’re three ______.", ["cup", "cups", "a cup", "cupes"], "B"),
        (33, "A: What ______ these?\nB: They’re books.", ["am", "is", "are", "be"], "C"),
        (34, "A: What ______ that?\nB: It’s a cup.", ["is", "are", "am", "be"], "A"),
        (35, "「這是什麼？」的正確英文是哪一個？", ["Who is this?", "Where is this?", "What is this?", "What are this?"], "C"),
        (36, "「那些是什麼？」的正確英文是哪一個？", ["What is those?", "What are those?", "What are that?", "Where are those?"], "B"),
        (37, "「它是一個盒子。」的正確英文是哪一個？", ["It’s a box.", "They’re boxes.", "It’s boxes.", "They’re a box."], "A"),
        (38, "「它們是我的手錶。」的正確英文是哪一個？", ["It’s my watch.", "They’re my watch.", "It’s my watches.", "They’re my watches."], "D"),
    ]),
    ("第五部分  詢問年齡", "請選出最適合的答案。", [
        (39, "How old ______ you?", ["am", "is", "are", "be"], "C"),
        (40, "How old ______ your brother?", ["is", "are", "am", "be"], "A"),
        (41, "How old ______ she?", ["am", "are", "be", "is"], "D"),
        (42, "A: How old are you?\nB: ______ twelve years old.", ["He’s", "I’m", "They’re", "She’s"], "B"),
        (43, "A: How old is Tom?\nB: ______ nine years old.", ["He’s", "I’m", "We’re", "They’re"], "A"),
    ]),
    ("第六部分  綜合練習", "請選出最適合的答案。", [
        (44, "dish 的複數是哪一個？", ["dishs", "dishes", "dish", "dishies"], "B"),
        (45, "bus 的複數是哪一個？", ["buss", "bus", "busies", "buses"], "D"),
        (46, "Those are four ______.", ["class", "classs", "classes", "a class"], "C"),
        (47, "She has ______ apple.", ["an", "a", "two", "these"], "A"),
        (48, "My sister has one ______.", ["watches", "watchs", "watching", "watch"], "D"),
        (49, "「那是什麼？」的正確英文是哪一個？", ["What are that?", "What is that?", "Where is that?", "Who is that?"], "B"),
        (50, "「這些是什麼？」的正確英文是哪一個？", ["What is this?", "Where are these?", "What are these?", "What is these?"], "C"),
    ]),
]


def set_font(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill="EAF0F6"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_question(doc, number, stem, options):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    lines = stem.split("\n")
    first = p.add_run(f"{number}. {lines[0]}")
    set_font(first, 10.5, bold=True)
    for line in lines[1:]:
        p.add_run().add_break()
        r = p.add_run(f"    {line}")
        set_font(r, 10.5)

    choices = doc.add_paragraph()
    choices.paragraph_format.left_indent = Cm(0.55)
    choices.paragraph_format.space_before = Pt(0)
    choices.paragraph_format.space_after = Pt(3)
    choices.paragraph_format.line_spacing = 1.0
    choices.paragraph_format.keep_together = True
    labels = ["A", "B", "C", "D"]
    for idx, (label, option) in enumerate(zip(labels, options)):
        if idx:
            choices.add_run("    ")
        r = choices.add_run(f"{label}. {option}")
        set_font(r, 10.5)


def add_section_heading(doc, title, instruction):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    keep_with_next(p)
    shade_paragraph(p)
    r = p.add_run(title)
    set_font(r, 12, bold=True)
    i = doc.add_paragraph()
    i.paragraph_format.space_before = Pt(0)
    i.paragraph_format.space_after = Pt(3)
    keep_with_next(i)
    r = i.add_run(instruction)
    set_font(r, 9.5, color="404040")


all_questions = [q for _, _, qs in SECTIONS for q in qs]
assert len(all_questions) == 50
assert [q[0] for q in all_questions] == list(range(1, 51))
assert len({q[1] for q in all_questions}) == 50
assert all(len(q[2]) == 4 for q in all_questions)
assert all(q[3] in "ABCD" for q in all_questions)

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.25)
section.bottom_margin = Cm(1.25)
section.left_margin = Cm(1.55)
section.right_margin = Cm(1.55)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
title_style.font.size = Pt(20)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(7)
title.add_run("英文基礎文法選擇題")

info = doc.add_table(rows=1, cols=3)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.autofit = False
widths = [Cm(7.2), Cm(7.2), Cm(3.0)]
labels = ["姓名：________________", "班級：________________", "分數：______"]
for cell, width, label in zip(info.rows[0].cells, widths, labels):
    cell.width = width
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=70, bottom=70, start=80, end=80)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label)
    set_font(r, 10.5)

intro = doc.add_paragraph()
intro.paragraph_format.space_before = Pt(5)
intro.paragraph_format.space_after = Pt(3)
r = intro.add_run("作答說明：本試卷共 50 題，每題有 A、B、C、D 四個選項，請選出一個最適合的答案。")
set_font(r, 10, bold=True)

# Manual breaks keep each learning skill together and make the paper easy to print.
break_before = {"第五部分  詢問年齡"}
for title_text, instruction, questions in SECTIONS:
    if title_text in break_before:
        doc.add_page_break()
    add_section_heading(doc, title_text, instruction)
    for number, stem, options, _ in questions:
        add_question(doc, number, stem, options)

doc.add_page_break()
answer_title = doc.add_paragraph(style="Title")
answer_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
answer_title.paragraph_format.space_after = Pt(8)
answer_title.add_run("教師用答案")

note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
note.paragraph_format.space_after = Pt(8)
r = note.add_run("每題 2 分，共 100 分")
set_font(r, 10.5)

answers = {number: answer for _, _, qs in SECTIONS for number, _, _, answer in qs}
table = doc.add_table(rows=6, cols=10)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = [str(i) for i in range(1, 11)]
for j, value in enumerate(headers):
    cell = table.cell(0, j)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=100, bottom=100)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "D9E2F3")
    cell._tc.get_or_add_tcPr().append(shade)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(value)
    set_font(r, 10, bold=True)

for row in range(1, 6):
    for col in range(10):
        number = row * 10 + col - 9
        cell = table.cell(row, col)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=110, bottom=110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{number}. {answers[number]}")
        set_font(r, 10.5, bold=True)

doc.core_properties.title = "英文基礎文法選擇題 50 題"
doc.core_properties.subject = "疑問詞 名詞複數 限定詞 What 問句與年齡問句"
doc.core_properties.author = "Ivan"
doc.save(OUTPUT)
print(OUTPUT)
